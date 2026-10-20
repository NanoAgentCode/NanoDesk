"""Generate document-chat fixtures with the bundled Office/PDF Python runtime."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas

root = Path(__file__).resolve().parents[2] / ".build" / "document-chat-qa"
root.mkdir(parents=True, exist_ok=True)
for extension in ["txt", "md", "csv"]:
    (root / f"中文 资料.{extension}").write_text("文档开头\n全文信息：项目代号 NanoDesk-QA，负责人小林。\n文档最后的验证码：TAIL-8972", encoding="utf-8")
(root / "utf16.txt").write_text("中文 UTF16 文档", encoding="utf-16")
(root / "gbk.txt").write_text("中文 GBK 文档", encoding="gbk")
(root / "empty.txt").write_text("  ", encoding="utf-8")
(root / "example.rtf").write_text(r"{\rtf1\ansi\uc1 Hello {\b world} \u20013?\u25991?\par TAIL-8972}", encoding="ascii")
doc = Document()
doc.add_paragraph("开头 A & B < C")
doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格文字 TABLE-8972"
doc.add_paragraph("最后一段 TAIL-8972")
doc.save(root / "example.docx")
book = Workbook()
book.active["A1"] = "第一张表"
book.active["B2"] = "TABLE-8972"
book.create_sheet("第二张表")["A2"] = "TAIL-8972"
book.save(root / "example.xlsx")
deck = Presentation()
for index in range(1, 12):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(100000, 100000, 3000000, 1000000).text = f"第 {index} 页 SLIDE-{index:02d} A & B"
deck.save(root / "example.pptx")
pdf = Canvas(str(root / "example.pdf"))
pdf.drawString(72, 720, "First page FIRST-8972")
pdf.showPage()
pdf.drawString(72, 720, "Last page TAIL-8972")
pdf.save()
scan = Canvas(str(root / "scan.pdf"))
scan.rect(50, 50, 100, 100)
scan.showPage()
scan.save()
print(root)
